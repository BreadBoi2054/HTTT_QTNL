# -*- coding: utf-8 -*-
"""
Master script to assemble the complete Major Assignment Report:
'Bao_Cao_Bai_Tap_Lon_HTTT_Quan_Tri_Nhan_Luc.docx'
"""

import os
import sys

# Configure UTF-8 for console output on Windows
sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START

# Ensure import paths
sys.path.insert(0, os.path.dirname(__file__))

from build_report import OUTPUT_FILE
from section_preamble import build_preamble
from section_intro import build_intro
from section_chapter1 import build_chapter1
from section_chapter2 import build_chapter2
from section_chapter3 import build_chapter3

def main():
    print("Initializing document...")
    doc = Document()

    # Configure page setup for Section 0 (A4 with Vietnamese standard margins)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)

    # Configure Header & Footer
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("Báo cáo BTL: Phân tích & Thiết kế Hệ thống Thông tin Quản trị Nhân lực (HRMIS)")
    hr.font.name = "Times New Roman"
    hr.font.size = Pt(9)
    hr.font.italic = True
    hr.font.color.rgb = RGBColor(148, 163, 184)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Học viện Hành chính và Quản trị công - 2026")
    fr.font.name = "Times New Roman"
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(148, 163, 184)

    # Configure default style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    print("Building Preamble (Title, Acknowledgements, Pledge, TOC, Lists)...")
    build_preamble(doc)

    print("Building Introduction (Phần Mở đầu)...")
    build_intro(doc)

    print("Building Chapter 1 (Cơ sở lý luận & Khảo sát bài toán)...")
    build_chapter1(doc)

    print("Building Chapter 2 (Phân tích & Thiết kế hệ thống)...")
    build_chapter2(doc)

    print("Building Chapter 3 & References (Kết quả, Đánh giá, Hướng phát triển)...")
    build_chapter3(doc)

    print(f"Saving final report document to: {OUTPUT_FILE}...")
    doc.save(OUTPUT_FILE)

    file_size = os.path.getsize(OUTPUT_FILE)
    print(f"Successfully created: {OUTPUT_FILE}")
    print(f"File size: {file_size / 1024 / 1024:.2f} MB")

if __name__ == '__main__':
    main()
